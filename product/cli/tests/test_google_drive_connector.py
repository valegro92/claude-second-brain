"""
Test del GoogleDriveConnector.

Strategia: iniettiamo un fake ``service`` (discovery client) tramite il
parametro ``service=`` del costruttore. Nessuna chiamata HTTP reale, nessun
OAuth flow.

Il fake service espone ``files()`` che ritorna un oggetto con ``list``,
``get_media``, ``export_media`` mockati per restituire i file desiderati.
"""

from __future__ import annotations

import io
import logging
from typing import Any
from unittest.mock import MagicMock

import pytest
from googleapiclient.errors import HttpError

from custodia_cli.connectors.google_drive import GoogleDriveConnector


# ---------------------------------------------------------------------------
# Fake Drive service
# ---------------------------------------------------------------------------


class _FakeList:
    """Mock di service.files().list(...).execute() ritorna i file forniti.

    Supporta sia il formato semplice ``{folder_id: [files...]}`` (single-page)
    sia il formato paginato ``{folder_id: [{"files": [...], "nextPageToken": "p2"}, ...]}``
    dove ogni elemento della lista esterna è una pagina.
    """

    def __init__(self, list_responses: dict[str, list[dict] | list[dict[str, Any]]]) -> None:
        # mapping parent_folder_id → lista file (semplice) o lista di pagine (paginated).
        self._responses = list_responses
        self.calls: list[dict] = []
        # tracker per pagination per-folder.
        self._page_index: dict[str, int] = {}

    def __call__(self, **kwargs: Any) -> "_FakeListRequest":
        self.calls.append(kwargs)
        # Estrai il folder_id dalla query "'<id>' in parents..."
        q = kwargs.get("q", "")
        folder_id = q.split("'")[1] if "'" in q else ""
        entry = self._responses.get(folder_id, [])

        # Formato paginato: lista di dict con chiave "files".
        if entry and isinstance(entry, list) and entry and isinstance(entry[0], dict) and "files" in entry[0]:
            idx = self._page_index.get(folder_id, 0)
            page = entry[idx] if idx < len(entry) else {"files": []}
            self._page_index[folder_id] = idx + 1
            return _FakeListRequest(dict(page))

        # Formato semplice: lista di file direttamente.
        return _FakeListRequest({"files": entry})


class _FakeListRequest:
    def __init__(self, payload: dict) -> None:
        self.payload = payload

    def execute(self) -> dict:
        return self.payload


class _FakeMediaRequest:
    """Mock del request object accettato da MediaIoBaseDownload."""

    def __init__(self, content: bytes, raise_after: int = 0) -> None:
        self.content = content
        self.raise_after = raise_after
        self.calls = 0


class _FakeGetRequest:
    """Mock di service.files().get(...).execute() per metadata folder."""

    def __init__(self, payload: dict) -> None:
        self.payload = payload

    def execute(self) -> dict:
        return self.payload


class _FakeFilesAPI:
    def __init__(
        self,
        list_responses: dict[str, list[dict] | list[dict[str, Any]]],
        media_bytes: dict[str, bytes],
        export_bytes: dict[tuple[str, str], bytes],
        download_errors: list[Exception] | None = None,
        folder_names: dict[str, str] | None = None,
    ) -> None:
        self.list = _FakeList(list_responses)
        self.media_bytes = media_bytes
        self.export_bytes = export_bytes
        self.download_errors = list(download_errors) if download_errors else []
        self.folder_names = folder_names or {}
        self.get_media_calls: list[dict[str, Any]] = []
        self.export_calls: list[tuple[str, str]] = []
        self.get_calls: list[dict[str, Any]] = []

    def get(self, **kwargs: Any) -> _FakeGetRequest:
        """Usata dal connector per recuperare il nome di una folder."""
        self.get_calls.append(kwargs)
        file_id = kwargs.get("fileId", "")
        return _FakeGetRequest(
            {"id": file_id, "name": self.folder_names.get(file_id, file_id)}
        )

    def get_media(self, **kwargs: Any) -> _FakeMediaRequest:
        self.get_media_calls.append(kwargs)
        file_id = kwargs.get("fileId", "")
        if self.download_errors:
            err = self.download_errors.pop(0)
            raise err
        return _FakeMediaRequest(self.media_bytes.get(file_id, b""))

    def export_media(self, *, fileId: str, mimeType: str) -> _FakeMediaRequest:
        self.export_calls.append((fileId, mimeType))
        return _FakeMediaRequest(self.export_bytes.get((fileId, mimeType), b""))


class _FakeService:
    def __init__(self, files_api: _FakeFilesAPI) -> None:
        self._files = files_api

    def files(self) -> _FakeFilesAPI:
        return self._files


# ---------------------------------------------------------------------------
# Helpers per generare bytes di file reali (riusa i fixture builder dei test parser)
# ---------------------------------------------------------------------------


def _make_simple_pdf_bytes() -> bytes:
    """Genera un PDF testuale minimale per test (usa pypdf low-level)."""
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (TestPDF Custodia) Tj ET")
    page[NameObject("/Contents")] = stream
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _make_simple_docx_bytes(text: str) -> bytes:
    import docx as docx_lib

    doc = docx_lib.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_simple_xlsx_bytes() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Foglio1"
    ws.append(["a", "b"])
    ws.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Helper: patch del downloader per usare i nostri bytes
# ---------------------------------------------------------------------------


@pytest.fixture
def patch_downloader(monkeypatch: pytest.MonkeyPatch):
    """Sostituisce MediaIoBaseDownload con una versione che legge da _FakeMediaRequest."""

    class _FakeDownloader:
        def __init__(self, buf: io.BytesIO, request: _FakeMediaRequest) -> None:
            self.buf = buf
            self.request = request
            self._done = False

        def next_chunk(self) -> tuple[Any, bool]:
            self.buf.write(self.request.content)
            self._done = True
            return (None, True)

    monkeypatch.setattr(
        "custodia_cli.connectors.google_drive.MediaIoBaseDownload",
        _FakeDownloader,
    )


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


def test_happy_path_mixed_files(patch_downloader: None) -> None:
    """4 file: PDF + GDoc + XLSX + video. Il connector produce 3 SourceDocument."""
    pdf_bytes = _make_simple_pdf_bytes()
    docx_export = _make_simple_docx_bytes("Documento Google esportato come DOCX")
    xlsx_bytes = _make_simple_xlsx_bytes()

    list_responses = {
        "root123": [
            {
                "id": "file_pdf",
                "name": "preventivo.pdf",
                "mimeType": "application/pdf",
                "size": "1024",
                "modifiedTime": "2026-01-01T10:00:00Z",
                "createdTime": "2026-01-01T09:00:00Z",
                "trashed": False,
                "parents": ["root123"],
                "md5Checksum": "abc123",
                "webViewLink": "https://drive.google.com/x",
            },
            {
                "id": "file_gdoc",
                "name": "verbale-meeting",
                "mimeType": "application/vnd.google-apps.document",
                "modifiedTime": "2026-02-02T10:00:00Z",
                "trashed": False,
                "parents": ["root123"],
            },
            {
                "id": "file_xlsx",
                "name": "listino.xlsx",
                "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "size": "2048",
                "modifiedTime": "2026-03-03T10:00:00Z",
                "trashed": False,
                "parents": ["root123"],
            },
            {
                "id": "file_video",
                "name": "demo.mp4",
                "mimeType": "video/mp4",
                "size": "10000000",
                "trashed": False,
                "parents": ["root123"],
            },
        ]
    }

    files_api = _FakeFilesAPI(
        list_responses=list_responses,
        media_bytes={"file_pdf": pdf_bytes, "file_xlsx": xlsx_bytes},
        export_bytes={
            (
                "file_gdoc",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ): docx_export
        },
    )
    service = _FakeService(files_api)
    connector = GoogleDriveConnector(root_folder_id="root123", service=service)

    docs = list(connector.iter_documents())

    assert len(docs) == 3
    by_id = {d.source_id: d for d in docs}

    # PDF
    assert "gdrive:file_pdf" in by_id
    pdf_doc = by_id["gdrive:file_pdf"]
    assert pdf_doc.mime_type == "application/pdf"
    assert "TestPDF" in pdf_doc.text
    assert pdf_doc.metadata["drive_id"] == "file_pdf"
    assert pdf_doc.metadata["modified_time"] == "2026-01-01T10:00:00Z"
    assert pdf_doc.metadata["web_view_link"] == "https://drive.google.com/x"

    # GDoc → esportato come DOCX
    gdoc = by_id["gdrive:file_gdoc"]
    assert "Documento Google" in gdoc.text
    assert (
        gdoc.mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    # XLSX
    xlsx = by_id["gdrive:file_xlsx"]
    assert "## Sheet: Foglio1" in xlsx.text

    # Stats
    stats = connector.stats
    assert stats["processed"] == 3
    assert stats["skipped_mime"] == 1  # il video


def test_trashed_files_are_skipped(patch_downloader: None) -> None:
    list_responses = {
        "root": [
            {
                "id": "trashed_one",
                "name": "cestino.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": True,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(list_responses, media_bytes={}, export_bytes={})
    connector = GoogleDriveConnector(root_folder_id="root", service=_FakeService(files_api))
    docs = list(connector.iter_documents())
    assert docs == []
    assert connector.stats["skipped_trashed"] == 1


def test_large_files_are_skipped_with_warning(
    caplog: pytest.LogCaptureFixture, patch_downloader: None
) -> None:
    too_big = 60 * 1024 * 1024  # 60 MB
    list_responses = {
        "root": [
            {
                "id": "huge",
                "name": "huge.pdf",
                "mimeType": "application/pdf",
                "size": str(too_big),
                "trashed": False,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(list_responses, media_bytes={}, export_bytes={})
    connector = GoogleDriveConnector(root_folder_id="root", service=_FakeService(files_api))

    with caplog.at_level(logging.WARNING, logger="custodia_cli.connectors.google_drive"):
        docs = list(connector.iter_documents())

    assert docs == []
    assert connector.stats["skipped_size"] == 1
    assert any("50MB" in r.message for r in caplog.records)


def test_dry_run_does_not_download(patch_downloader: None) -> None:
    """dry_run=True: il connector produce SourceDocument con text='' senza chiamare get_media."""
    list_responses = {
        "root": [
            {
                "id": "f1",
                "name": "x.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
                "modifiedTime": "2026-01-01T00:00:00Z",
            }
        ]
    }
    files_api = _FakeFilesAPI(list_responses, media_bytes={}, export_bytes={})
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api), dry_run=True
    )

    docs = list(connector.iter_documents())

    assert len(docs) == 1
    assert docs[0].text == ""
    assert docs[0].source_id == "gdrive:f1"
    assert files_api.get_media_calls == []  # nessun download
    assert files_api.export_calls == []


def test_retry_on_http_429_then_success(
    patch_downloader: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Una HttpError 429 sul download → retry → success al 2° tentativo."""
    # Stub time.sleep nel modulo del connector per non rallentare i test.
    monkeypatch.setattr("custodia_cli.connectors.google_drive.time.sleep", lambda _s: None)

    # Costruiamo un HttpError 429 manualmente.
    fake_resp = MagicMock()
    fake_resp.status = 429
    fake_resp.reason = "Too Many Requests"
    http_error = HttpError(resp=fake_resp, content=b"rate limit")

    pdf_bytes = _make_simple_pdf_bytes()

    list_responses = {
        "root": [
            {
                "id": "f1",
                "name": "p.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"f1": pdf_bytes},
        export_bytes={},
        download_errors=[http_error],  # solo il 1° call fallisce
    )
    connector = GoogleDriveConnector(root_folder_id="root", service=_FakeService(files_api))

    docs = list(connector.iter_documents())
    assert len(docs) == 1
    # get_media chiamato 2 volte (1 fallito + 1 success)
    assert len(files_api.get_media_calls) == 2


def test_recursive_folder_traversal(patch_downloader: None) -> None:
    """Folder con sottofolder → BFS attraversa entrambe."""
    pdf_bytes = _make_simple_pdf_bytes()
    list_responses = {
        "root": [
            {
                "id": "subfolder",
                "name": "sub",
                "mimeType": "application/vnd.google-apps.folder",
                "trashed": False,
                "parents": ["root"],
            },
            {
                "id": "top_pdf",
                "name": "top.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            },
        ],
        "subfolder": [
            {
                "id": "nested_pdf",
                "name": "nested.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["subfolder"],
            }
        ],
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"top_pdf": pdf_bytes, "nested_pdf": pdf_bytes},
        export_bytes={},
    )
    connector = GoogleDriveConnector(root_folder_id="root", service=_FakeService(files_api))
    docs = list(connector.iter_documents())
    ids = {d.source_id for d in docs}
    assert ids == {"gdrive:top_pdf", "gdrive:nested_pdf"}


def test_google_slides_are_skipped(patch_downloader: None) -> None:
    list_responses = {
        "root": [
            {
                "id": "slides1",
                "name": "deck",
                "mimeType": "application/vnd.google-apps.presentation",
                "trashed": False,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(list_responses, media_bytes={}, export_bytes={})
    connector = GoogleDriveConnector(root_folder_id="root", service=_FakeService(files_api))
    docs = list(connector.iter_documents())
    assert docs == []
    assert connector.stats["skipped_mime"] == 1


# ---------------------------------------------------------------------------
# Nuovi test: hardening dopo review (FIX A1, A2, A3, A9, A10, A11, A12)
# ---------------------------------------------------------------------------


def test_trashed_folder_is_not_recursed(patch_downloader: None) -> None:
    """FIX A1: una folder con ``trashed=True`` non deve essere accodata al BFS.

    Anche se la query Drive filtra ``trashed=false``, il connector deve fare
    una seconda verifica client-side: se per qualsiasi motivo (race, cache API)
    una folder trashed ritornasse, non vogliamo scendere dentro.
    """
    list_responses = {
        "root": [
            {
                "id": "trashed_folder",
                "name": "cestino",
                "mimeType": "application/vnd.google-apps.folder",
                "trashed": True,
                "parents": ["root"],
            }
        ],
        "trashed_folder": [
            {
                "id": "should_not_appear",
                "name": "ghost.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["trashed_folder"],
            }
        ],
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"should_not_appear": b"xxx"},
        export_bytes={},
        folder_names={"root": "Root"},
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    docs = list(connector.iter_documents())
    assert docs == []
    # Trashed folder contata come skipped_trashed.
    assert connector.stats["skipped_trashed"] >= 1
    # E la lista NON deve essere chiamata su trashed_folder (no BFS recursion).
    folder_ids_listed = {
        call.get("q", "").split("'")[1] if "'" in call.get("q", "") else ""
        for call in files_api.list.calls
    }
    assert "trashed_folder" not in folder_ids_listed


def test_list_query_filters_trashed_false(patch_downloader: None) -> None:
    """FIX A1: la query ``q`` deve contenere ``and trashed=false``."""
    list_responses: dict[str, list[dict]] = {"root": []}
    files_api = _FakeFilesAPI(
        list_responses, media_bytes={}, export_bytes={}, folder_names={"root": "Root"}
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    list(connector.iter_documents())
    assert files_api.list.calls, "files.list deve essere chiamato almeno una volta"
    for call in files_api.list.calls:
        assert "trashed=false" in call.get("q", "")


def test_per_file_unexpected_exception_does_not_kill_scan(
    patch_downloader: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """FIX A2: un errore non-HttpError/ParserError (es. ConnectionResetError)
    su un singolo file deve essere catturato e contato in ``errors``, senza
    interrompere lo scan degli altri file."""
    pdf_bytes = _make_simple_pdf_bytes()

    list_responses = {
        "root": [
            {
                "id": "broken",
                "name": "broken.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            },
            {
                "id": "ok",
                "name": "ok.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            },
        ]
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"ok": pdf_bytes},
        export_bytes={},
        folder_names={"root": "Root"},
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )

    # Patch _download_file: solleva ConnectionResetError SOLO per "broken".
    original_download = connector._download_file

    def _fake_download(service: Any, file_id: str) -> bytes:
        if file_id == "broken":
            raise ConnectionResetError("connessione chiusa dal server")
        return original_download(service, file_id)

    monkeypatch.setattr(connector, "_download_file", _fake_download)

    docs = list(connector.iter_documents())
    # Il file ok deve essere processato; il broken contato in errors.
    ids = {d.source_id for d in docs}
    assert "gdrive:ok" in ids
    assert "gdrive:broken" not in ids
    assert connector.stats["errors"] == 1
    assert connector.stats["processed"] == 1


def test_source_path_preserves_drive_hierarchy(patch_downloader: None) -> None:
    """FIX A3: ``source_path`` deve riflettere la gerarchia di folder Drive.

    Struttura:
        Root/
          Sub1/
            Sub2/
              nested.pdf
    Atteso: ``source_path = "/Root/Sub1/Sub2/nested.pdf"``.
    """
    pdf_bytes = _make_simple_pdf_bytes()
    list_responses = {
        "root": [
            {
                "id": "sub1",
                "name": "Sub1",
                "mimeType": "application/vnd.google-apps.folder",
                "trashed": False,
                "parents": ["root"],
            }
        ],
        "sub1": [
            {
                "id": "sub2",
                "name": "Sub2",
                "mimeType": "application/vnd.google-apps.folder",
                "trashed": False,
                "parents": ["sub1"],
            }
        ],
        "sub2": [
            {
                "id": "leaf",
                "name": "nested.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["sub2"],
            }
        ],
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"leaf": pdf_bytes},
        export_bytes={},
        folder_names={"root": "Root"},
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    docs = list(connector.iter_documents())
    assert len(docs) == 1
    assert docs[0].source_path == "/Root/Sub1/Sub2/nested.pdf"
    # Verifica anche che files.get(root) sia stato chiamato per recuperare il nome.
    assert any(call.get("fileId") == "root" for call in files_api.get_calls)


def test_pagination_merges_multiple_pages(patch_downloader: None) -> None:
    """FIX A9: ``files.list`` con ``nextPageToken`` deve essere ri-chiamato
    con ``pageToken=...`` finché non si esauriscono le pagine.
    """
    pdf_bytes = _make_simple_pdf_bytes()
    # 2 pagine: la prima ha nextPageToken="p2", la seconda no.
    page_files = [
        {
            "id": f"f{i}",
            "name": f"file{i}.pdf",
            "mimeType": "application/pdf",
            "size": "100",
            "trashed": False,
            "parents": ["root"],
        }
        for i in range(3)
    ]
    page2_files = [
        {
            "id": f"g{i}",
            "name": f"file{i}.pdf",
            "mimeType": "application/pdf",
            "size": "100",
            "trashed": False,
            "parents": ["root"],
        }
        for i in range(2)
    ]
    list_responses: dict[str, list[dict[str, Any]]] = {
        "root": [
            {"files": page_files, "nextPageToken": "p2"},
            {"files": page2_files},
        ]
    }
    media = {**{f"f{i}": pdf_bytes for i in range(3)}, **{f"g{i}": pdf_bytes for i in range(2)}}
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes=media,
        export_bytes={},
        folder_names={"root": "Root"},
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    docs = list(connector.iter_documents())
    assert len(docs) == 5
    # Almeno 2 chiamate list(): la prima senza pageToken, la seconda con pageToken="p2".
    assert len(files_api.list.calls) >= 2
    tokens = [c.get("pageToken") for c in files_api.list.calls]
    assert "p2" in tokens
    # Primo call: pageToken None.
    assert files_api.list.calls[0].get("pageToken") is None


@pytest.mark.parametrize(
    "status,reason,error_details,should_retry",
    [
        (500, "Internal Server Error", None, True),
        (503, "Service Unavailable", None, True),
        (429, "Too Many Requests", None, True),
        (403, "Forbidden", [{"reason": "userRateLimitExceeded"}], True),
        (403, "Forbidden", [{"reason": "rateLimitExceeded"}], True),
        (403, "Forbidden", [{"reason": "insufficientPermissions"}], False),
    ],
)
def test_retry_decision_by_status_and_reason(
    status: int,
    reason: str,
    error_details: list[dict] | None,
    should_retry: bool,
) -> None:
    """FIX A10: ``_is_retriable_http_error`` discrimina su status + reason
    structured, NON su substring match del messaggio.
    """
    from custodia_cli.connectors.google_drive import _is_retriable_http_error

    fake_resp = MagicMock()
    fake_resp.status = status
    fake_resp.reason = reason
    err = HttpError(resp=fake_resp, content=b"x")
    # Inietta error_details in modo esplicito (googleapiclient ≥2.5 lo espone).
    if error_details is not None:
        err.error_details = error_details  # type: ignore[attr-defined]

    assert _is_retriable_http_error(err) is should_retry


def test_403_insufficient_permissions_does_not_retry(
    patch_downloader: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    """FIX A10 (integration): un 403 con reason ``insufficientPermissions``
    si propaga immediatamente, senza retry loop. Lo scan continua sugli altri file.
    """
    monkeypatch.setattr(
        "custodia_cli.connectors.google_drive.time.sleep", lambda _s: None
    )
    fake_resp = MagicMock()
    fake_resp.status = 403
    fake_resp.reason = "Forbidden"
    perm_error = HttpError(resp=fake_resp, content=b"no perms")
    perm_error.error_details = [{"reason": "insufficientPermissions"}]  # type: ignore[attr-defined]

    pdf_bytes = _make_simple_pdf_bytes()
    list_responses = {
        "root": [
            {
                "id": "denied",
                "name": "denied.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            },
            {
                "id": "ok",
                "name": "ok.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            },
        ]
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"ok": pdf_bytes},
        export_bytes={},
        download_errors=[perm_error],  # solo il primo download fallisce
        folder_names={"root": "Root"},
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    docs = list(connector.iter_documents())
    # Il primo file fallisce (no retry) ma è catturato come HttpError → errors+=1.
    # Il secondo passa.
    ids = {d.source_id for d in docs}
    assert ids == {"gdrive:ok"}
    assert connector.stats["errors"] == 1
    # Verifica che NON ci siano stati retry: 1 sola chiamata su "denied".
    denied_calls = [c for c in files_api.get_media_calls if c.get("fileId") == "denied"]
    assert len(denied_calls) == 1


def test_list_call_uses_shared_drives_flags(patch_downloader: None) -> None:
    """FIX A11: ``files.list`` deve essere chiamato con
    ``supportsAllDrives=True`` e ``includeItemsFromAllDrives=True``.
    """
    list_responses: dict[str, list[dict]] = {"root": []}
    files_api = _FakeFilesAPI(
        list_responses, media_bytes={}, export_bytes={}, folder_names={"root": "Root"}
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    list(connector.iter_documents())
    assert files_api.list.calls
    for call in files_api.list.calls:
        assert call.get("supportsAllDrives") is True
        assert call.get("includeItemsFromAllDrives") is True


def test_get_media_uses_supports_all_drives(patch_downloader: None) -> None:
    """FIX A11: ``files.get_media`` deve passare ``supportsAllDrives=True``."""
    pdf_bytes = _make_simple_pdf_bytes()
    list_responses = {
        "root": [
            {
                "id": "f1",
                "name": "f1.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"f1": pdf_bytes},
        export_bytes={},
        folder_names={"root": "Root"},
    )
    connector = GoogleDriveConnector(
        root_folder_id="root", service=_FakeService(files_api)
    )
    list(connector.iter_documents())
    assert files_api.get_media_calls
    for call in files_api.get_media_calls:
        assert call.get("supportsAllDrives") is True


def test_cache_dir_writes_raw_bytes(patch_downloader: None, tmp_path) -> None:
    """FIX A12: con ``cache_dir`` settato, i raw bytes scaricati vengono
    persistiti su disco con nome basato sul ``source_id`` sanitizzato.
    """
    from custodia_cli.connectors.google_drive import _sanitize_for_path

    pdf_bytes = _make_simple_pdf_bytes()
    list_responses = {
        "root": [
            {
                "id": "f1",
                "name": "preventivo.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"f1": pdf_bytes},
        export_bytes={},
        folder_names={"root": "Root"},
    )
    cache_dir = tmp_path / "cache"
    connector = GoogleDriveConnector(
        root_folder_id="root",
        service=_FakeService(files_api),
        cache_dir=cache_dir,
    )
    docs = list(connector.iter_documents())
    assert len(docs) == 1

    expected_path = cache_dir / _sanitize_for_path("gdrive:f1")
    assert expected_path.exists()
    assert expected_path.read_bytes() == pdf_bytes
    # Verifica metadata: cached_at popolato.
    assert docs[0].metadata.get("cached_at") == str(expected_path)


def test_cache_dir_skips_rewrite_if_exists(
    patch_downloader: None, tmp_path
) -> None:
    """FIX A4: se il blob di cache esiste già, non riscriviamo (evita IO inutile)."""
    from custodia_cli.connectors.google_drive import _sanitize_for_path

    pdf_bytes = _make_simple_pdf_bytes()
    list_responses = {
        "root": [
            {
                "id": "f1",
                "name": "preventivo.pdf",
                "mimeType": "application/pdf",
                "size": "100",
                "trashed": False,
                "parents": ["root"],
            }
        ]
    }
    files_api = _FakeFilesAPI(
        list_responses,
        media_bytes={"f1": pdf_bytes},
        export_bytes={},
        folder_names={"root": "Root"},
    )
    cache_dir = tmp_path / "cache"
    cache_dir.mkdir()
    # Pre-popola la cache con un blob diverso (mtime "vecchio").
    cache_path = cache_dir / _sanitize_for_path("gdrive:f1")
    cache_path.write_bytes(b"stale-cached-bytes")
    original_mtime = cache_path.stat().st_mtime_ns

    connector = GoogleDriveConnector(
        root_folder_id="root",
        service=_FakeService(files_api),
        cache_dir=cache_dir,
    )
    list(connector.iter_documents())
    # Il file su disco non è stato ri-scritto.
    assert cache_path.read_bytes() == b"stale-cached-bytes"
    assert cache_path.stat().st_mtime_ns == original_mtime
